"""GCB Asia 2020 parser — XLSX codebook backbone + DOCX corroboration.

The official codebook (210114_GCB10_Asia_Codebook_Final_AFG.xlsx) is a
mechanically clean VARIABLE/QUESTION/VALUES/VALUE LABELS grid: it supplies
every variable (including battery sub-items Q1A..) with full question text
and per-code labels. The master questionnaire DOCX supplies response-scale
tables we use to corroborate the XLSX parse: a variable whose DOCX scale
codes match its XLSX codes gets higher parse_confidence.

Wave key: "asia2020" (matches the gcb spec's source: key).
"""

from __future__ import annotations

import re
from pathlib import Path

import docx
import openpyxl

from ..crossvalidate.polarity import classify_label
from ..paths import repo_root
from ..schema import CodebookRow

WAVE = "asia2020"
_ORIGINALS = "data/gcb/questionnaires/originals"
XLSX_REL = f"{_ORIGINALS}/210114_GCB10_Asia_Codebook_Final_AFG.xlsx"
DOCX_REL = f"{_ORIGINALS}/GCB_2020_Asia_Master_Questionnaire_Final.docx"

_HEADER = {"VARIABLE", "QUESTION", "VALUES", "VALUE LABELS"}
_DO_NOT_READ = re.compile(r"^\s*do not read\s*[:\-]?\s*", re.IGNORECASE)
_VAR_PARA = re.compile(r"^([A-Z][A-Z0-9_]{0,15})\.\s*$")


def _is_missing_label(label: str) -> bool:
    if _DO_NOT_READ.match(label):
        return True
    return classify_label(label) == "missing"


def _parse_xlsx(path: Path) -> dict[str, tuple[str, dict[str, str]]]:
    """var -> (question_text, {code: label}) from the codebook grid."""
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    col_idx: dict[str, int] = {}
    out: dict[str, tuple[str, dict[str, str]]] = {}
    var: str | None = None
    q_parts: list[str] = []
    scale: dict[str, str] = {}

    def flush() -> None:
        nonlocal var, q_parts, scale
        if var is not None and scale:
            out[var] = (" ".join(q_parts).strip(), scale)
        var, q_parts, scale = None, [], {}

    for row in ws.iter_rows(values_only=True):
        cells = ["" if v is None else str(v).strip() for v in row]
        if not col_idx:
            uppers = {c.upper() for c in cells if c}
            if _HEADER <= uppers:
                col_idx = {c.upper(): i for i, c in enumerate(cells) if c}
            continue
        v_cell = cells[col_idx["VARIABLE"]]
        q_cell = cells[col_idx["QUESTION"]]
        code_cell = cells[col_idx["VALUES"]]
        label_cell = cells[col_idx["VALUE LABELS"]]
        if v_cell:
            flush()
            var = v_cell
        if var is None:
            continue
        if q_cell:
            q_parts.append(q_cell)
        if code_cell != "" and label_cell != "":
            # Normalize float-formatted integer codes ("1.0" -> "1").
            code = code_cell
            try:
                f = float(code)
                code = str(int(f)) if f == int(f) else str(f)
            except ValueError:
                pass
            scale[code] = label_cell
    flush()
    wb.close()
    return out


def _iter_block_items(doc: docx.document.Document) -> list[tuple[str, object]]:
    """Paragraphs and tables in true body order."""
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    items: list[tuple[str, object]] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            items.append(("p", Paragraph(child, doc)))
        elif child.tag == qn("w:tbl"):
            items.append(("tbl", Table(child, doc)))
    return items


def _parse_docx_scales(path: Path) -> dict[str, dict[str, str]]:
    """var -> {code: label} from the questionnaire's response tables.

    A variable paragraph looks like "Q2." / "GEN."; the next 2-column table
    whose second column is numeric is taken as its response scale. FILTER
    note tables (1x1) are skipped.
    """
    doc = docx.Document(str(path))
    items = _iter_block_items(doc)
    out: dict[str, dict[str, str]] = {}
    current: str | None = None
    for kind, obj in items:
        if kind == "p":
            m = _VAR_PARA.match(obj.text.strip())  # type: ignore[union-attr]
            if m:
                current = m.group(1)
            continue
        # kind == "tbl"
        tbl = obj
        rows = tbl.rows  # type: ignore[union-attr]
        if current is None or len(rows) < 2:
            continue
        scale: dict[str, str] = {}
        for r in rows:
            cells = [c.text.strip() for c in r.cells]
            if len(cells) < 2:
                continue
            label, code = cells[0], cells[-1]
            if re.fullmatch(r"-?[0-9]+", code) and label:
                scale[code] = label
        if len(scale) >= 2:
            out.setdefault(current, scale)
            current = None
    return out


def parse_gcb() -> list[CodebookRow]:
    root = repo_root()
    xlsx_path = root / XLSX_REL
    docx_path = root / DOCX_REL
    xlsx = _parse_xlsx(xlsx_path)
    docx_scales = _parse_docx_scales(docx_path) if docx_path.is_file() else {}

    def docx_match(var: str) -> dict[str, str] | None:
        # Dataset names diverge from questionnaire codes systematically:
        # TQ2 <- Q2 (T prefix), Q30_Asia <- Q30 (regional suffix).
        candidates = [var]
        if var.startswith("T"):
            candidates.append(var[1:])
        if var.endswith("_Asia"):
            candidates.append(var[: -len("_Asia")])
        for cand in candidates:
            if cand in docx_scales:
                return docx_scales[cand]
        return None

    rows: list[CodebookRow] = []
    for var, (question, scale) in xlsx.items():
        docx_scale = docx_match(var)
        corroborated = docx_scale is not None and set(docx_scale) >= set(scale)
        confidence = 0.95 if corroborated else 0.85
        method = "xlsx_codebook+docx" if corroborated else "xlsx_codebook"
        for code, label in scale.items():
            clean = _DO_NOT_READ.sub("", label).strip()
            rows.append(
                CodebookRow(
                    survey="gcb",
                    wave=WAVE,
                    raw_var=var,
                    question_text=question,
                    response_code=code,
                    response_label=clean,
                    missing_code_flag=_is_missing_label(label),
                    source_doc=XLSX_REL,
                    source_page="Sheet1",
                    language="en",
                    notes=None,
                    parse_confidence=confidence,
                    parse_method=method,
                )
            )
    return rows
